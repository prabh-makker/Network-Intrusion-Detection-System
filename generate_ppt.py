from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation():
    prs = Presentation()
    
    # Slide 1: Title Slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "NIDS Sentinel"
    subtitle.text = "AI-Powered Real-Time Network Intrusion Detection & Threat Analysis\n\nProject Presentation"

    # Slide 2: Introduction
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    title.text = "Introduction"
    tf = body_shape.text_frame
    tf.text = "What is NIDS Sentinel?"
    p = tf.add_paragraph()
    p.text = "A robust Network Intrusion Detection System designed to monitor network traffic for malicious activity."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Leverages advanced Artificial Intelligence and Machine Learning algorithms."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Provides real-time visibility into network health and security threats."
    p.level = 1

    # Slide 3: Key Features
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    title.text = "Key Features"
    tf = body_shape.text_frame
    tf.text = "Real-Time Packet Sniffing: Live traffic analysis using Scapy."
    p = tf.add_paragraph()
    p.text = "AI-Powered Classification: High accuracy detection of DoS, DDoS, Probe, and U2R attacks."
    p = tf.add_paragraph()
    p.text = "Live Dashboard: Interactive React-based dashboard with real-time charts."
    p = tf.add_paragraph()
    p.text = "Threat Alerts & AI Explainability: Understand why a packet was flagged as a threat."

    # Slide 4: System Architecture
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    title.text = "System Architecture"
    tf = body_shape.text_frame
    tf.text = "Modular and Scalable Design:"
    p = tf.add_paragraph()
    p.text = "Packet Sniffer (Scapy): Captures raw network packets."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "ML Inference Engine: Extracts features and classifies packets."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "FastAPI Backend: Handles data processing, database interactions, and WebSockets."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Next.js Frontend: Displays live data to security analysts."
    p.level = 1

    # Slide 5: Technology Stack
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    title.text = "Technology Stack"
    tf = body_shape.text_frame
    tf.text = "Frontend: Next.js 16, React, TypeScript, Tailwind CSS, Recharts"
    p = tf.add_paragraph()
    p.text = "Backend: FastAPI (Python), SQLAlchemy, WebSockets"
    p = tf.add_paragraph()
    p.text = "Database: SQLite (Ready for PostgreSQL)"
    p = tf.add_paragraph()
    p.text = "Machine Learning: Scikit-learn (Random Forest), Joblib"
    p = tf.add_paragraph()
    p.text = "Networking/Sniffing: Scapy"

    # Slide 6: Machine Learning Model
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    title.text = "Machine Learning Details"
    tf = body_shape.text_frame
    tf.text = "Algorithm: Random Forest Classifier"
    p = tf.add_paragraph()
    p.text = "Trained on KDD-Cup 99/NSL-KDD datasets (over 25,000 packets)."
    p = tf.add_paragraph()
    p.text = "Accuracy: Achieves ~99.82% detection accuracy."
    p = tf.add_paragraph()
    p.text = "Features Evaluated: 12 key network features including protocol type, flags, byte counts, and error rates."

    # Slide 7: Threat Detection in Action
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    title.text = "Threat Detection in Action"
    tf = body_shape.text_frame
    tf.text = "Continuous Monitoring: Dashboard updates instantly as traffic flows."
    p = tf.add_paragraph()
    p.text = "Automated Alerting: Critical threats are flagged automatically."
    p = tf.add_paragraph()
    p.text = "Explainable AI: The system highlights which features (e.g., unusual byte count) triggered the alert."
    p = tf.add_paragraph()
    p.text = "Actionable Insights: Analysts can view severity, mitigation strategies, and take action."

    # Slide 8: Future Enhancements
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    title.text = "Future Enhancements"
    tf = body_shape.text_frame
    tf.text = "Transition from Intrusion Detection to Prevention (IPS capabilities)."
    p = tf.add_paragraph()
    p.text = "Implement Deep Learning models for complex temporal attack patterns."
    p = tf.add_paragraph()
    p.text = "Cloud integrations and distributed sensor networks for enterprise-scale deployments."
    p = tf.add_paragraph()
    p.text = "Enhanced SIEM integration and automated reporting."

    # Slide 9: Conclusion
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    title.text = "Conclusion"
    tf = body_shape.text_frame
    tf.text = "NIDS Sentinel offers a state-of-the-art solution for network security."
    p = tf.add_paragraph()
    p.text = "Bridges the gap between raw packet data and actionable human insights."
    p = tf.add_paragraph()
    p.text = "Empowers organizations to defend against evolving cyber threats proactively."

    # Slide 10: Q&A
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Questions?"
    subtitle.text = "Thank You for Your Attention"

    prs.save('NIDS_Sentinel_Presentation.pptx')
    print("Successfully created NIDS_Sentinel_Presentation.pptx")

if __name__ == '__main__':
    create_presentation()
