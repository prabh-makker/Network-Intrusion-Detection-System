from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))

def add_box(slide, text, left, top, width, height, fill_color_hex, font_size=16):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = hex_to_rgb(fill_color_hex)
    line = shape.line
    line.color.rgb = hex_to_rgb("FFFFFF")
    line.width = Pt(1.5)
    
    text_frame = shape.text_frame
    text_frame.text = text
    text_frame.word_wrap = True
    
    p = text_frame.paragraphs[0]
    p.font.size = Pt(font_size)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb("FFFFFF")
    p.alignment = PP_ALIGN.CENTER
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shape

def add_arrow(slide, left, top, width, height, rotation=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
    shape.rotation = rotation
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = hex_to_rgb("CBD5E1")
    line = shape.line
    line.color.rgb = hex_to_rgb("FFFFFF")
    line.width = Pt(1)
    return shape

def apply_dark_theme(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = hex_to_rgb("0F172A") # Slate 900

def create_presentation():
    prs = Presentation()
    
    # -----------------------------
    # SLIDE 1: Title Slide
    # -----------------------------
    title_slide_layout = prs.slide_layouts[6] # Blank
    slide1 = prs.slides.add_slide(title_slide_layout)
    apply_dark_theme(slide1)
    
    txBox = slide1.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "NIDS Sentinel"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb("38BDF8")
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "Architecture & Data Flow"
    p2.font.size = Pt(32)
    p2.font.color.rgb = hex_to_rgb("F8FAFC")
    p2.alignment = PP_ALIGN.CENTER

    # -----------------------------
    # SLIDE 2: Flowchart
    # -----------------------------
    blank_slide_layout = prs.slide_layouts[6]
    slide2 = prs.slides.add_slide(blank_slide_layout)
    apply_dark_theme(slide2)
    
    # Title
    txBox = slide2.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "System Architecture Data Flow"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb("FFFFFF")
    p.alignment = PP_ALIGN.CENTER

    # Flowchart dimensions
    box_w = Inches(2.8)
    box_h = Inches(1.0)
    center_x = Inches(3.6) 
    
    # Layer 1: Network Traffic
    box1_y = Inches(1.2)
    add_box(slide2, "🌐 Raw Network Traffic\n(Packets)", center_x, box1_y, box_w, box_h, "2563EB") # Blue
    
    # Arrow 1
    add_arrow(slide2, center_x + box_w/2 - Inches(0.15), box1_y + box_h, Inches(0.3), Inches(0.4))
    
    # Layer 2: Packet Sniffer
    box2_y = box1_y + box_h + Inches(0.4)
    add_box(slide2, "🔍 Packet Sniffer\n(Scapy & PyPcap)", center_x, box2_y, box_w, box_h, "7C3AED") # Purple
    
    # Arrow 2
    add_arrow(slide2, center_x + box_w/2 - Inches(0.15), box2_y + box_h, Inches(0.3), Inches(0.4))

    # Layer 3: Feature Extraction & ML
    box3_y = box2_y + box_h + Inches(0.4)
    add_box(slide2, "🧠 ML Inference Engine\n(Random Forest & Joblib)", center_x, box3_y, box_w, Inches(1.2), "DB2777") # Pink
    
    # Arrow 3
    add_arrow(slide2, center_x + box_w/2 - Inches(0.15), box3_y + Inches(1.2), Inches(0.3), Inches(0.4))

    # Layer 4: Backend
    box4_y = box3_y + Inches(1.2) + Inches(0.4)
    add_box(slide2, "⚙️ FastAPI Backend\n(Data Processing & WebSockets)", center_x, box4_y, box_w, box_h, "059669") # Green
    
    # Output Arrows
    # Left arrow
    left_arrow = slide2.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, center_x + Inches(0.5), box4_y + box_h, Inches(0.3), Inches(0.5))
    left_arrow.rotation = 45
    left_arrow.fill.solid()
    left_arrow.fill.fore_color.rgb = hex_to_rgb("CBD5E1")
    
    # Right arrow
    right_arrow = slide2.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, center_x + box_w - Inches(0.8), box4_y + box_h, Inches(0.3), Inches(0.5))
    right_arrow.rotation = -45
    right_arrow.fill.solid()
    right_arrow.fill.fore_color.rgb = hex_to_rgb("CBD5E1")

    # Layer 5: Databases & Dashboards
    box5_y = box4_y + box_h + Inches(0.5)
    
    # Left Box (DB)
    add_box(slide2, "🗄️ Threat Database\n(SQLite / PostgreSQL)", Inches(1.0), box5_y, Inches(2.6), box_h, "D97706") # Orange
    
    # Right Box (Dashboard)
    add_box(slide2, "📊 Live React Dashboard\n(Next.js & Recharts)", Inches(6.4), box5_y, Inches(2.6), box_h, "D97706") # Orange

    # Alerts Box
    alert_y = box3_y
    alert_x = Inches(7.0)
    add_box(slide2, "🚨 Threat Alerts\n(Block / Mitigation)", alert_x, alert_y, Inches(2.0), box_h, "DC2626", font_size=14) # Red
    
    # Arrow from ML to Alerts
    arrow_right = slide2.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, center_x + box_w + Inches(0.1), alert_y + box_h/2 - Inches(0.15), Inches(0.5), Inches(0.3))
    arrow_right.fill.solid()
    arrow_right.fill.fore_color.rgb = hex_to_rgb("CBD5E1")
    
    # -----------------------------
    # SLIDE 3: Component Details
    # -----------------------------
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide3 = prs.slides.add_slide(slide_layout)
    apply_dark_theme(slide3)
    
    title = slide3.shapes.title
    title.text = "Architecture Details"
    title.text_frame.paragraphs[0].font.color.rgb = hex_to_rgb("FFFFFF")
    
    body_shape = slide3.shapes.placeholders[1]
    tf = body_shape.text_frame
    
    # Setting text and formatting
    p = tf.paragraphs[0]
    p.text = "1. Packet Sniffer (Scapy): Captures real-time raw packets on the network interface."
    p.font.color.rgb = hex_to_rgb("E2E8F0")
    
    p2 = tf.add_paragraph()
    p2.text = "2. Feature Extraction: Extracts 12 critical features (e.g., protocol, flags, byte count)."
    p2.font.color.rgb = hex_to_rgb("E2E8F0")
    
    p3 = tf.add_paragraph()
    p3.text = "3. ML Inference: Random Forest categorizes into Normal, DoS, DDoS, Probe, or U2R."
    p3.font.color.rgb = hex_to_rgb("E2E8F0")
    
    p4 = tf.add_paragraph()
    p4.text = "4. FastAPI Backend: Orchestrates data flow and broadcasts results via WebSockets."
    p4.font.color.rgb = hex_to_rgb("E2E8F0")
    
    p5 = tf.add_paragraph()
    p5.text = "5. Live Dashboard: React UI visualizes the traffic and threats instantly."
    p5.font.color.rgb = hex_to_rgb("E2E8F0")

    prs.save('NIDS_Decorative_Flowchart.pptx')
    print("Successfully created NIDS_Decorative_Flowchart.pptx")

if __name__ == '__main__':
    create_presentation()
